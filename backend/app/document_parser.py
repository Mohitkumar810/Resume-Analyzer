from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class DocumentParsingError(ValueError):
    """Raised when a supported document cannot be read or has no text."""


def extract_document_text(filename: str, content: bytes) -> str:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension == "pdf":
        text = _extract_pdf_text(content)
    elif extension == "docx":
        text = _extract_docx_text(content)
    else:
        raise DocumentParsingError("Upload a PDF or DOCX resume.")

    cleaned_text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not cleaned_text:
        raise DocumentParsingError(
            "No readable text was found. Scanned PDFs need OCR before they can be analyzed."
        )
    return cleaned_text


def _extract_pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise DocumentParsingError("Password-protected PDFs are not supported.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError, OSError) as exc:
        if isinstance(exc, DocumentParsingError):
            raise
        raise DocumentParsingError("The PDF could not be read. Check that it is a valid PDF.") from exc


def _extract_docx_text(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError, OSError) as exc:
        raise DocumentParsingError(
            "The DOCX could not be read. Check that it is a valid Word document."
        ) from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    blocks.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(blocks)
